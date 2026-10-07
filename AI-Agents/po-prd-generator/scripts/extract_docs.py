#!/usr/bin/env python3
"""Extrae texto de documentos (pdf, docx, xlsx, pptx, md, txt, csv, json, html)
a archivos .txt UTF-8 (+ .outline.txt con títulos y números de línea) en una
carpeta de trabajo, para que el agente los lea por tramos sin cargar el binario
ni imprimir el contenido en consola.

Uso: python extract_docs.py <archivo> [<archivo> ...] [--out DIR]
Salida (stdout): una línea por documento -> ruta_txt | chars | unidades | outline.
"""
import argparse
import re
import sys
import tempfile
from pathlib import Path

PLAIN = {".md", ".txt", ".csv", ".json", ".html", ".htm", ".xml", ".yml", ".yaml"}


OUTLINE = re.compile(
    r"^(#{1,4} \S|=== (Página|Hoja|Slide)"
    r"|\d{1,2}(\.\d{1,2}){1,3}\.?\s+[A-ZÁÉÍÓÚÑ]"
    r"|\d{1,2}\.\s+[A-ZÁÉÍÓÚÑ][^.*?¿:]{2,70}$"
    r"|[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ0-9 ,/()-]{4,}$)")


def clean(text: str) -> str:
    return text.encode("utf-8", "replace").decode("utf-8", "replace")


def outline(text: str) -> str:
    """Líneas de título/página/hoja con su número de línea, para leer por offset."""
    rows = [f"{n}: {line.strip()[:100]}" for n, line in enumerate(text.splitlines(), 1)
            if len(line.strip()) <= 120 and OUTLINE.match(line.strip())]
    return "\n".join(rows)


def pdf_text(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            body = page.extract_text() or ""
        except Exception as exc:  # página corrupta: seguir
            body = f"[error de extracción: {exc}]"
        parts.append(f"=== Página {i} ===\n{body}")
    return "\n\n".join(parts), f"{len(reader.pages)} págs"


def docx_text(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            txt = p.text.strip()
            if not txt:
                continue
            m = re.match(r"(?:Heading|Título)\s*(\d)", p.style.name or "")
            out.append(("#" * int(m.group(1)) + " " + txt) if m else txt)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out), f"{len(doc.paragraphs)} párrafos, {len(doc.tables)} tablas"


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Hoja: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            cells = ["" if c is None else str(c).strip().replace("\n", " ") for c in row]
            out.append("| " + " | ".join(cells).rstrip(" |") + " |")
    return "\n".join(out), f"{len(wb.worksheets)} hojas"


def pptx_text(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"=== Slide {i} ===")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out), f"{len(prs.slides)} slides"


def plain_text(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return text, f"{text.count(chr(10)) + 1} líneas"


HANDLERS = {".pdf": pdf_text, ".docx": docx_text, ".xlsx": xlsx_text,
            ".xlsm": xlsx_text, ".pptx": pptx_text}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", help="carpeta de salida (default: temp del sistema)")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    out_dir = Path(args.out) if args.out else Path(tempfile.gettempdir()) / "po-prd-generator"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = 0
    for name in args.files:
        src = Path(name)
        ext = src.suffix.lower()
        if not src.exists():
            print(f"ERROR | {name} | no existe")
            rc = 1
            continue
        handler = HANDLERS.get(ext) or (plain_text if ext in PLAIN else None)
        if handler is None:
            print(f"ERROR | {name} | formato no soportado ({ext})")
            rc = 1
            continue
        try:
            text, units = handler(src)
        except ImportError as exc:
            print(f"ERROR | {name} | falta dependencia: {exc.name} (pip install {exc.name})")
            rc = 1
            continue
        except Exception as exc:
            print(f"ERROR | {name} | {type(exc).__name__}: {exc}")
            rc = 1
            continue
        text = clean(text)
        dest = out_dir / f"{src.stem}.txt"
        dest.write_text(text, encoding="utf-8")
        toc = outline(text)
        (out_dir / f"{src.stem}.outline.txt").write_text(toc, encoding="utf-8")
        print(f"{dest} | {len(text)} chars | {units} | outline: {toc.count(chr(10)) + bool(toc)} líneas")
    return rc


if __name__ == "__main__":
    sys.exit(main())
